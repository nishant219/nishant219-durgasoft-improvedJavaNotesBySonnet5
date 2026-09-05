#!/usr/bin/env python3
"""Recover Markdown from the existing .docx / .doc notes.

.docx is preferred (tables sit in document order). For notes that only exist
as .doc, tables were appended at the end of the body by the old converter --
they are emitted under a `## Tables (placement lost)` heading so the rewrite
pass can put them back where they belong.
"""

from __future__ import annotations

import html
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
NOTES = ROOT / "notes"


def md_table(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * cols]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return out + [""]


# --------------------------------------------------------------------- docx


def run_is_code(run) -> bool:
    return (run.font.name or "") in ("Consolas", "Courier New")


def para_to_md(p: Paragraph) -> list[str]:
    style = p.style.name if p.style is not None else ""
    text = p.text or ""
    if not text.strip():
        return []

    if style.startswith("Heading"):
        level = int("".join(c for c in style if c.isdigit()) or 2)
        return ["#" * level + " " + text.strip(), ""]

    runs = [r for r in p.runs if r.text]
    if runs and all(run_is_code(r) for r in runs):
        return ["```java", *text.split("\n"), "```", ""]

    # inline marks
    parts = []
    for r in runs:
        t = r.text
        if not t.strip():
            parts.append(t)
        elif run_is_code(r):
            parts.append(f"`{t}`")
        elif r.bold:
            parts.append(f"**{t}**")
        elif r.italic:
            parts.append(f"*{t}*")
        else:
            parts.append(t)
    line = "".join(parts) or text

    if "List Bullet" in style:
        return ["- " + line.strip()]
    if "List Number" in style:
        return ["1. " + line.strip()]
    if p.paragraph_format.left_indent and all(r.italic for r in runs if r.text.strip()):
        return ["> " + line.strip(), ""]
    return [line.strip(), ""]


def docx_to_md(src: Path) -> str:
    doc = Document(src)
    body = doc.element.body
    para_map = {p._element: p for p in doc.paragraphs}
    table_map = {t._element: t for t in doc.tables}

    out: list[str] = []
    for child in body.iterchildren():
        if child in para_map:
            out += para_to_md(para_map[child])
        elif child in table_map:
            t: Table = table_map[child]
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
            out += md_table(rows)
    return collapse("\n".join(out))


# ---------------------------------------------------------------------- doc


class DocHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.tables: list[list[list[str]]] = []
        self.buf: list[str] = []
        self.tag: str | None = None
        self.in_body = False
        self.row: list[str] | None = None
        self.rows: list[list[str]] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "body":
            self.in_body = True
            return
        if not self.in_body:
            return
        if tag == "table":
            self.rows = []
        elif tag == "tr":
            self.row = []
        elif tag in ("td", "th", "h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote"):
            self.tag = tag
            self.buf = []

    def handle_data(self, data):
        if self.in_body and self.tag:
            self.buf.append(data)

    def handle_endtag(self, tag):
        if tag == "body":
            self.in_body = False
            return
        if not self.in_body:
            return
        text = "".join(self.buf).strip()
        if tag in ("td", "th"):
            if self.row is not None:
                self.row.append(text.replace("\n", " "))
        elif tag == "tr":
            if self.rows is not None and self.row:
                self.rows.append(self.row)
            self.row = None
        elif tag == "table":
            if self.rows:
                self.tables.append(self.rows)
            self.rows = None
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.out += ["#" * int(tag[1]) + " " + text, ""]
        elif tag == "pre":
            self.out += ["```java", *"".join(self.buf).strip("\n").split("\n"), "```", ""]
        elif tag == "blockquote":
            self.out += ["> " + text, ""]
        elif tag == "li":
            self.out.append("- " + text)
        elif tag == "p":
            if text:
                self.out += [text, ""]
        if tag in ("td", "th", "h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote"):
            self.tag = None
            self.buf = []


def doc_to_md(src: Path) -> str:
    p = DocHTML()
    p.feed(src.read_text(encoding="utf-8"))
    out = list(p.out)
    if p.tables:
        out += ["", "## Tables (placement lost -- re-place these in context)", ""]
        for rows in p.tables:
            out += md_table(rows)
    return collapse("\n".join(out))


def collapse(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def recover(stem_dir: Path, stem: str) -> str | None:
    docx = stem_dir / f"{stem}.docx"
    doc = stem_dir / f"{stem}.doc"
    if docx.exists():
        return docx_to_md(docx)
    if doc.exists():
        return doc_to_md(doc)
    return None


def main() -> None:
    if len(sys.argv) >= 2:
        src = Path(sys.argv[1])
        text = docx_to_md(src) if src.suffix == ".docx" else doc_to_md(src)
        sys.stdout.write(text)
        return

    n = 0
    for d in sorted(NOTES.glob("chunk-*")):
        stems = sorted({f.stem for f in d.iterdir() if f.suffix in (".doc", ".docx")})
        for stem in stems:
            dest = d / f"{stem}.md"
            if dest.exists():
                continue
            text = recover(d, stem)
            if text is None:
                continue
            dest.write_text(text, encoding="utf-8")
            n += 1
            print(f"recovered {dest.relative_to(ROOT)}")
    print(f"{n} markdown files recovered")


if __name__ == "__main__":
    main()
