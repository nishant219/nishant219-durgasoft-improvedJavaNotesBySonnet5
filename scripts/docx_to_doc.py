#!/usr/bin/env python3
"""Write Word-compatible .doc (HTML) files from .docx notes."""

from __future__ import annotations

import html
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
NOTES = ROOT / "notes"

HEAD = """<html xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:w="urn:schemas-microsoft-com:office:word">
<head>
<meta charset="utf-8">
<meta http-equiv="Content-Type" content="text/html; charset=utf-8">
<title>{title}</title>
<style>
body {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; line-height: 1.35; }}
h1 {{ font-size: 18pt; }}
h2 {{ font-size: 14pt; }}
h3 {{ font-size: 12pt; }}
pre, .code {{ font-family: Consolas, Courier New, monospace; font-size: 9pt;
  background: #f4f4f4; padding: 8px; white-space: pre-wrap; }}
table {{ border-collapse: collapse; margin: 8px 0 16px 0; }}
td, th {{ border: 1px solid #666; padding: 4px 8px; vertical-align: top; }}
th {{ font-weight: bold; }}
blockquote {{ margin-left: 16px; font-style: italic; color: #333; }}
</style>
</head>
<body>
"""
TAIL = "</body></html>\n"


def looks_like_code(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if t.startswith(("class ", "public ", "package ", "import ", "//", "/*", "{", "}")):
        return True
    if t.endswith("{") or t.endswith(";") or t.endswith("}"):
        if any(k in t for k in ("int ", "void ", "new ", "String", "boolean ", "byte ", "System.out", "return ", "if (", "for (")):
            return True
    return False


def convert_docx(src: Path, dest: Path) -> None:
    d = Document(src)
    parts = [HEAD.format(title=html.escape(src.stem))]
    in_code = False
    code_lines: list[str] = []

    def flush_code() -> None:
        nonlocal in_code, code_lines
        if in_code:
            parts.append("<pre class='code'>" + html.escape("\n".join(code_lines)) + "</pre>\n")
            in_code = False
            code_lines = []

    for p in d.paragraphs:
        style = p.style.name if p.style is not None else ""
        text = p.text or ""
        if style.startswith("Heading"):
            flush_code()
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append(f"<h{level}>{html.escape(text)}</h{level}>\n")
            continue
        if "List Bullet" in style:
            flush_code()
            parts.append(f"<ul><li>{html.escape(text)}</li></ul>\n")
            continue
        if "List Number" in style:
            flush_code()
            parts.append(f"<ol><li>{html.escape(text)}</li></ol>\n")
            continue
        if looks_like_code(text) or (in_code and text.startswith(("    ", "\t", "//", "}", "{"))):
            in_code = True
            code_lines.append(text)
            continue
        flush_code()
        if not text.strip():
            continue
        parts.append(f"<p>{html.escape(text)}</p>\n")

    flush_code()

    for table in d.tables:
        parts.append("<table>\n")
        for i, row in enumerate(table.rows):
            parts.append("<tr>")
            tag = "th" if i == 0 else "td"
            for cell in row.cells:
                parts.append(f"<{tag}>{html.escape(cell.text)}</{tag}>")
            parts.append("</tr>\n")
        parts.append("</table>\n")

    parts.append(TAIL)
    dest.write_text("".join(parts), encoding="utf-8")


def main() -> None:
    if len(sys.argv) >= 3:
        convert_docx(Path(sys.argv[1]), Path(sys.argv[2]))
        print(sys.argv[2])
        return
    n = 0
    for src in sorted(NOTES.rglob("*.docx")):
        dest = src.with_suffix(".doc")
        convert_docx(src, dest)
        print(f"wrote {dest.relative_to(ROOT)}")
        n += 1
    print(f"converted {n} docx → doc")


if __name__ == "__main__":
    main()
