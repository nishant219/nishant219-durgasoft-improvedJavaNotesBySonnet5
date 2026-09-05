#!/usr/bin/env python3
"""Convert our lecture-note Markdown files to Word .docx."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
NOTES = ROOT / "notes"

HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
FENCE = re.compile(r"^```")
UL = re.compile(r"^[-*]\s+(.*)$")
OL = re.compile(r"^(\d+)\.\s+(.*)$")
TABLE_SEP = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
HR = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD = re.compile(r"\*\*(.+?)\*\*")
ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
CODE = re.compile(r"`([^`]+)`")


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_runs(paragraph, text: str, *, code: bool = False) -> None:
    if code:
        run = paragraph.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        return

    pos = 0
    token = re.compile(
        r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*(?!\*)(?:(?!\*).)+?(?<!\*)\*(?!\*))"
    )
    for m in token.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        raw = m.group(0)
        if raw.startswith("**") and raw.endswith("**"):
            run = paragraph.add_run(raw[2:-2])
            run.bold = True
        elif raw.startswith("`") and raw.endswith("`"):
            run = paragraph.add_run(raw[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif raw.startswith("["):
            lm = LINK.match(raw)
            run = paragraph.add_run(lm.group(1) if lm else raw)
            run.font.color.rgb = RGBColor(0x0B, 0x57, 0xD0)
            run.underline = True
        elif raw.startswith("*") and raw.endswith("*"):
            run = paragraph.add_run(raw[1:-1])
            run.italic = True
        else:
            paragraph.add_run(raw)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, row[j] if j < len(row) else "")
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def convert_text(md: str, dest: Path, title: str | None = None) -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    if title:
        doc.core_properties.title = title

    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]
        if FENCE.match(line.strip()):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                r = run._element
                rPr = r.get_or_add_rPr()
                # keep as preformatted paragraph
                code_buf = []
                in_code = False
            else:
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if HR.match(stripped):
            i += 1
            continue

        hm = HEADING.match(stripped)
        if hm:
            level = min(len(hm.group(1)), 4)
            text = hm.group(2).strip()
            p = doc.add_heading(text, level=level)
            i += 1
            continue

        if stripped.startswith(">"):
            quote = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            add_runs(p, quote)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1].strip()):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            add_table(doc, rows)
            continue

        um = UL.match(stripped)
        if um:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, um.group(1))
            i += 1
            continue
        om = OL.match(stripped)
        if om:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, om.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def convert_file(src: Path, dest: Path) -> None:
    convert_text(src.read_text(encoding="utf-8"), dest, title=src.stem)


def main() -> None:
    if len(sys.argv) >= 3:
        convert_file(Path(sys.argv[1]), Path(sys.argv[2]))
        print(sys.argv[2])
        return
    count = 0
    for md in sorted(NOTES.rglob("*.md")):
        dest = md.with_suffix(".docx")
        convert_file(md, dest)
        print(f"wrote {dest.relative_to(ROOT)}")
        count += 1
    print(f"converted {count} files")


if __name__ == "__main__":
    main()
